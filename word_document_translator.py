import requests
import re
import html
import os
import sys
from docx import Document
from docx.text.paragraph import Paragraph

OLLAMA_URL = "http://localhost:11434"
MODEL_CONTEXT = "qwen2.5:7b"

LANGUAGE_SOURCE = "english"
LANGUAGE_TARGET = "french"

REQUEST_TIMEOUT_SECONDS = 60

INPUT_DOCX_PATH = "original_document.docx"
OUTPUT_DOCX_PATH = "translated_document.docx"


# ---------------------------------------------------------------------------
# 1. OLLAMA COMMUNICATION FUNCTIONS
# ---------------------------------------------------------------------------

class OllamaUnavailableError(Exception):
    """Raised when the Ollama service cannot be reached."""


class OllamaQueryError(Exception):
    """Raised when a request to Ollama fails (timeout, HTTP error, etc.)."""

def build_ollama_endpoint(ollama_url, endpoint):
    """Build a stable Ollama endpoint from a base URL or legacy full endpoint URL."""
    base_url = ollama_url.rstrip("/")
    for suffix in ("/api/generate", "/api/tags"):
        if base_url.endswith(suffix):
            base_url = base_url[:-len(suffix)]
            break
    return f"{base_url}{endpoint}"

def check_ollama_reachable(ollama_url=OLLAMA_URL, timeout=5):
    """Checks if the Ollama service is reachable before translation starts."""
    health_url = build_ollama_endpoint(ollama_url, "/api/tags")
    try:
        response = requests.get(health_url, timeout=timeout)
        response.raise_for_status()
    except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as exc:
        raise OllamaUnavailableError(
            f"Ollama is unreachable at {health_url}. "
            f"Check that the Ollama service is running and accessible."
        ) from exc
    except requests.exceptions.RequestException as exc:
        raise OllamaUnavailableError(
            f"Ollama health check failed ({health_url}): {exc}"
        ) from exc

def generate_document_context(full_text, ollama_url=OLLAMA_URL, model_context=MODEL_CONTEXT, source_language=LANGUAGE_SOURCE, target_language=LANGUAGE_TARGET):
    """Analyzes the entire document to create a universal context summary."""
    url = build_ollama_endpoint(ollama_url, "/api/generate")
    prompt = (
        f"Analyze the following text extracted from a document. "
        f"Generate a brief summary (2-3 sentences max) describing only: "
        f"1. The main subject of the document.\n"
        f"2. The tone to adopt (formal, literary, technical, legal, marketing, etc.).\n"
        f"3. The target audience.\n"
        f"Be extremely concise and objective.\n\nDocument text:\n{full_text}"
    )
    payload = {
        "model": model_context,
        "prompt": prompt,
        "stream": False,
        "options": {
            "temperature": 0.3
        }
    }
    try:
        response = requests.post(url, json=payload, timeout=REQUEST_TIMEOUT_SECONDS)
        response.raise_for_status()
        return response.json().get("response", "").strip()
    except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as e:
        print(f"Error during context generation (Ollama query failed): {e}")
        raise OllamaQueryError("Context generation failed because Ollama query did not complete.") from e
    except requests.exceptions.RequestException as e:
        print(f"Error during context generation (HTTP/API failure): {e}")
        raise OllamaQueryError("Context generation failed due to an Ollama API error.") from e
    except ValueError as e:
        print(f"Error during context generation (invalid JSON response): {e}")
        return "No specific context."

def translate_text_with_context(tagged_text, context, ollama_url=OLLAMA_URL, model_context=MODEL_CONTEXT, source_language=LANGUAGE_SOURCE, target_language=LANGUAGE_TARGET):
    """Translates the text generically while filtering out AI conversational preambles."""
    # Anti-hallucination security: avoid calling the API if the text has no letters
    if not re.search(r'[a-zA-Z]', tagged_text):
        return tagged_text

    url = build_ollama_endpoint(ollama_url, "/api/generate")
    
    # Restructured prompt to "force" direct translation start
    prompt = f"""You are an expert translator.
    Global context of the document: "{context}"

    ABSOLUTE AND IMPERATIVE RULES:
    1. Translate the ENTIRE text FROM {source_language} into {target_language}. Leave absolutely NO word, part of a sentence, or start of a sentence in {source_language}. EVERYTHING must be translated.
    2. NEVER SPECIFY that the text is a "Direct {target_language} translation" or a "Literal {target_language} translation" if the text is very generic or lacks specific context.
    3. NEVER add comments or preambles for text that is too small or empty to be translated. Just translate it as is or leave it empty.
    4. Pay close attention to quotes (" or '): you must translate everything before, inside, and after the quotes.
    5. Keep EXACTLY the position of HTML tags (<...>) without translating their code.
    6. NEVER copy the source text in {source_language} into your response.
    7. FORMAL PROHIBITION to add introductory or polite words.    

    Original text:
    {tagged_text}

    Direct {target_language} translation:"""

    payload = {
        "model": model_context,
        "prompt": prompt,
        "stream": False,
        "options": {
            "temperature": 0.3
        }
    }
    
    try:
        response = requests.post(url, json=payload, timeout=REQUEST_TIMEOUT_SECONDS)
        response.raise_for_status()
        translated_text = response.json().get("response", "").strip()
        
        # --- POST-TRANSLATION CLEANUP (Anti-preamble) ---
        # Remove typical introductory sentences if the model disobeys
        patterns_to_remove = [
            # French patterns
            r"^(Voici|Voilà) la traduction(.*?):\s*",
            r"^Voici le texte traduit(.*?):\s*",
            r"^Traduction(.*?):\s*",
            r"^Bien sûr(.*?):\s*",
            r"^Je suis désolé(.*?)\n",
            # English patterns
            r"^(Here is|Here's) the translation(.*?):\s*",
            r"^(Here is|Here's) the translated text(.*?):\s*",
            r"^Translation(.*?):\s*",
            r"^(Sure|Of course)(.*?):\s*",
            r"^I am sorry(.*?)\n"
        ]
        
        for pattern in patterns_to_remove:
            translated_text = re.sub(pattern, "", translated_text, flags=re.IGNORECASE | re.DOTALL).strip()
            
        return translated_text

    except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as e:
        print(f"Error during translation (Ollama query failed): {e}")
        return tagged_text
    except requests.exceptions.RequestException as e:
        print(f"Error during translation (HTTP/API failure): {e}")
        return tagged_text
    except ValueError as e:
        print(f"Error during translation (invalid JSON response): {e}")
        return tagged_text

def free_ollama_memory(ollama_url=OLLAMA_URL, model=MODEL_CONTEXT):
    """Forces Ollama to unload the model from memory (VRAM/RAM)."""
    payload = {
        "model": model,
        "keep_alive": 0
    }
    try:
        requests.post(build_ollama_endpoint(ollama_url, "/api/generate"), json=payload)
        print(f" -> Memory successfully freed for the model {model}.")
    except Exception as e:
        print(f" -> Error freeing memory: {e}")


class TranslationProgress:
    """Simple terminal progress bar for translation completion."""
    def __init__(self, total_items, bar_width=32):
        self.total_items = max(0, total_items)
        self.current_items = 0
        self.bar_width = bar_width

    def render(self):
        if self.total_items <= 0:
            return

        ratio = min(1.0, self.current_items / self.total_items)
        filled = int(self.bar_width * ratio)
        percent = int(ratio * 100)
        bar = "#" * filled + "-" * (self.bar_width - filled)
        sys.stdout.write(
            f"\r   Progress: [{bar}] {self.current_items}/{self.total_items} ({percent:3d}%)"
        )
        sys.stdout.flush()

    def advance(self, step=1):
        if self.total_items <= 0:
            return
        self.current_items = min(self.total_items, self.current_items + step)
        self.render()

    def finish(self):
        if self.total_items <= 0:
            return
        self.current_items = self.total_items
        self.render()
        print()

# ---------------------------------------------------------------------------
# 2. TAG MANIPULATION AND PYTHON-DOCX LOGIC
# ---------------------------------------------------------------------------

def paragraph_to_tags(paragraph):
    """Transforms a paragraph of runs into a tagged text string (HTML-like)."""
    tagged_text = ""
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        # Escape potential existing tags to avoid XML conflicts
        text = text.replace("<", "&lt;").replace(">", "&gt;")
        
        # Apply tags according to the run's style
        if run.bold: text = f"<b>{text}</b>"
        if run.italic: text = f"<i>{text}</i>"
        if run.underline: text = f"<u>{text}</u>"
        if run.font.strike: text = f"<s>{text}</s>"
        
        tagged_text += text
    return tagged_text


def is_translatable_text(tagged_text):
    """Determine whether a tagged text should be sent for translation."""
    return bool(tagged_text and tagged_text.strip() and re.search(r'[a-zA-Z]', tagged_text))

def snapshot_run_style(run):
    """Capture a run style so it can be reapplied after translation."""
    style = {
        "bold": run.bold,
        "italic": run.italic,
        "underline": run.underline,
        "strike": run.font.strike,
        "name": run.font.name,
        "size": run.font.size,
        "highlight_color": run.font.highlight_color,
        "style": run.style,
        "color_rgb": None,
        "color_theme": None,
    }

    color = run.font.color
    if color is not None:
        try:
            style["color_rgb"] = color.rgb
        except Exception:
            style["color_rgb"] = None

        if style["color_rgb"] is None:
            try:
                style["color_theme"] = color.theme_color
            except Exception:
                style["color_theme"] = None

    return style


def apply_style_snapshot(style, run):
    """Reapply a style snapshot to a run."""
    run.bold = style["bold"]
    run.italic = style["italic"]
    run.underline = style["underline"]
    run.font.strike = style["strike"]
    run.font.name = style["name"]

    if style["size"] is not None:
        run.font.size = style["size"]

    if style["highlight_color"] is not None:
        run.font.highlight_color = style["highlight_color"]

    if style["style"] is not None:
        run.style = style["style"]

    if style["color_rgb"] is not None:
        try:
            run.font.color.rgb = style["color_rgb"]
            return
        except Exception:
            pass

    if style["color_theme"] is not None:
        try:
            run.font.color.theme_color = style["color_theme"]
        except Exception:
            pass


def build_paragraph_style_plan(paragraph):
    """Build a style plan from original runs based on text lengths."""
    style_plan = []
    for run in paragraph.runs:
        run_text = run.text or ""
        if not run_text:
            continue
        style_plan.append({
            "length": len(run_text),
            "style": snapshot_run_style(run),
        })

    if not style_plan and paragraph.runs:
        style_plan.append({
            "length": 1,
            "style": snapshot_run_style(paragraph.runs[0]),
        })

    return style_plan


def apply_tagged_translation_preserving_style(paragraph, translated_text):
    """Parses translated text and rebuilds runs while preserving original run styles."""
    if not paragraph.runs:
        return

    style_plan = build_paragraph_style_plan(paragraph)
    if not style_plan:
        return

    # Clean up potential markdown blocks hallucinated by the LLM
    translated_text = re.sub(r'```[a-zA-Z0-9]*', '', translated_text).strip()
    
    # Empty existing runs
    p_element = paragraph._p
    for run in paragraph.runs:
        p_element.remove(run._r)
        
    # Global regex that captures any HTML tag
    tag_pattern = r'(</?\w+(?:\s+[^>]+)?>)'
    tokens = re.split(tag_pattern, translated_text)
    
    bold_level = 0
    italic_level = 0
    underline_level = 0
    strike_level = 0

    style_index = 0
    style_remaining = style_plan[style_index]["length"]
    
    for token in tokens:
        if not token:
            continue
            
        # If the token is an HTML tag
        if token.startswith("<") and token.endswith(">"):
            match_name = re.match(r'</?(\w+)', token)
            if not match_name:
                continue
            tag_name = match_name.group(1).lower()
            is_closing = token.startswith("</")
            
            # Handle styles and common aliases (strong/bold, em/italic)
            if tag_name in ["b", "strong"]:
                bold_level += -1 if is_closing else 1
            elif tag_name in ["i", "em"]:
                italic_level += -1 if is_closing else 1
            elif tag_name in ["u"]:
                underline_level += -1 if is_closing else 1
            elif tag_name in ["s", "strike", "del"]:
                strike_level += -1 if is_closing else 1
            
            # Safeguard to prevent going below 0
            bold_level = max(0, bold_level)
            italic_level = max(0, italic_level)
            underline_level = max(0, underline_level)
            strike_level = max(0, strike_level)
            
            # Other tags (<span>, <div>, etc.) are silently ignored.
            
        else:
            # It's raw text: convert HTML entities (e.g., &amp; -> &)
            clean_text = html.unescape(token)

            while clean_text:
                if style_remaining <= 0 and style_index < len(style_plan) - 1:
                    style_index += 1
                    style_remaining = style_plan[style_index]["length"]

                # Keep last style for overflow when translation is longer than source.
                if style_remaining <= 0:
                    style_remaining = len(clean_text)

                chunk_len = min(len(clean_text), style_remaining)
                text_chunk = clean_text[:chunk_len]
                clean_text = clean_text[chunk_len:]

                run = paragraph.add_run(text_chunk)
                apply_style_snapshot(style_plan[style_index]["style"], run)

                # Override preserved run style only if translated tags request emphasis.
                if bold_level > 0:
                    run.bold = True
                if italic_level > 0:
                    run.italic = True
                if underline_level > 0:
                    run.underline = True
                if strike_level > 0:
                    run.font.strike = True

                style_remaining -= chunk_len


def iter_header_footer_parts(doc):
    """Yield each distinct header/footer part once across all sections."""
    seen_partnames = set()
    for section in doc.sections:
        candidates = [
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ]
        for part in candidates:
            try:
                partname = str(part.part.partname)
            except Exception:
                partname = None

            if partname and partname in seen_partnames:
                continue
            if partname:
                seen_partnames.add(partname)

            yield part


def extract_text_from_story_part(story_part, texts):
    """Collect paragraphs, tables, and text boxes text from a body/header/footer part."""
    for p in story_part.paragraphs:
        if p.text.strip():
            texts.append(p.text)

    for table in story_part.tables:
        extract_text_from_table(table, texts)

    for txbx in story_part._element.xpath('.//w:txbxContent'):
        for p_element in txbx.xpath('.//w:p'):
            p = Paragraph(p_element, story_part)
            if p.text.strip():
                texts.append(p.text)


def extract_text_from_table(table, texts):
    """Recursively collect text from a table and nested tables."""
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.strip():
                    texts.append(p.text)
            for nested_table in cell.tables:
                extract_text_from_table(nested_table, texts)


def count_translatable_items_in_table(table):
    """Count translatable paragraph items recursively in a table."""
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                tagged_text = paragraph_to_tags(paragraph)
                if is_translatable_text(tagged_text):
                    count += 1
            for nested_table in cell.tables:
                count += count_translatable_items_in_table(nested_table)
    return count


def count_translatable_items_in_story_part(story_part):
    """Count all translatable paragraph items in a body/header/footer part."""
    count = 0

    for paragraph in story_part.paragraphs:
        tagged_text = paragraph_to_tags(paragraph)
        if is_translatable_text(tagged_text):
            count += 1

    for table in story_part.tables:
        count += count_translatable_items_in_table(table)

    for txbx in story_part._element.xpath('.//w:txbxContent'):
        for p_element in txbx.xpath('.//w:p'):
            paragraph = Paragraph(p_element, story_part)
            tagged_text = paragraph_to_tags(paragraph)
            if is_translatable_text(tagged_text):
                count += 1

    return count


def translate_paragraph_if_needed(paragraph, context, ollama_url, model_context, source_language, target_language, progress=None):
    """Translate one paragraph when text content requires translation."""
    tagged_text = paragraph_to_tags(paragraph)
    if is_translatable_text(tagged_text):
        translation = translate_text_with_context(
            tagged_text,
            context,
            ollama_url=ollama_url,
            model_context=model_context,
            source_language=source_language,
            target_language=target_language,
        )
        apply_tagged_translation_preserving_style(paragraph, translation)
        if progress is not None:
            progress.advance()


def translate_table(table, context, ollama_url, model_context, source_language, target_language, progress=None):
    """Translate a table recursively, including nested tables."""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                translate_paragraph_if_needed(
                    paragraph,
                    context,
                    ollama_url,
                    model_context,
                    source_language,
                    target_language,
                    progress=progress,
                )
            for nested_table in cell.tables:
                translate_table(
                    nested_table,
                    context,
                    ollama_url,
                    model_context,
                    source_language,
                    target_language,
                    progress=progress,
                )


def translate_story_part(story_part, context, ollama_url, model_context, source_language, target_language, progress=None):
    """Translate all translatable text in a body/header/footer part."""
    for paragraph in story_part.paragraphs:
        translate_paragraph_if_needed(
            paragraph,
            context,
            ollama_url,
            model_context,
            source_language,
            target_language,
            progress=progress,
        )

    for table in story_part.tables:
        translate_table(
            table,
            context,
            ollama_url,
            model_context,
            source_language,
            target_language,
            progress=progress,
        )

    for txbx in story_part._element.xpath('.//w:txbxContent'):
        for p_element in txbx.xpath('.//w:p'):
            paragraph = Paragraph(p_element, story_part)
            translate_paragraph_if_needed(
                paragraph,
                context,
                ollama_url,
                model_context,
                source_language,
                target_language,
                progress=progress,
            )

# ---------------------------------------------------------------------------
# 3. MAIN PROCESS
# ---------------------------------------------------------------------------

def extract_all_text(input_path):
    """Scans the document a first time to extract all raw text."""
    doc = Document(input_path)
    texts = []

    # Body text (paragraphs, tables, text boxes)
    extract_text_from_story_part(doc, texts)

    # Headers and footers across sections
    for story_part in iter_header_footer_parts(doc):
        extract_text_from_story_part(story_part, texts)
            
    # Returns the first 8000 characters to give a general idea without crashing the API
    return "\n".join(texts)[:8000]

def translate_full_docx(input_path, output_path, ollama_url, model_context, source_language, target_language):
    if not os.path.isfile(input_path):
        print(f"Error: input file not found: {input_path}")
        return False

    try:
        check_ollama_reachable(ollama_url=ollama_url)
    except OllamaUnavailableError as e:
        print(f"Error: {e}")
        return False

    print("1. Extracting text for context analysis...")
    try:
        full_text = extract_all_text(input_path)
    except FileNotFoundError:
        print(f"Error: input file not found: {input_path}")
        return False
    except Exception as e:
        print(f"Error while reading input document: {e}")
        return False
    
    print(f"2. Generating context with AI model ({model_context})...")
    try:
        context = generate_document_context(full_text, ollama_url=ollama_url, model_context=model_context, source_language=source_language, target_language=target_language)
    except OllamaQueryError as e:
        print(f"Warning: {e}")
        context = "No specific context."
    print(f"   [Retained context] : {context}\n")
    
    try:
        doc = Document(input_path)
    except FileNotFoundError:
        print(f"Error: input file not found: {input_path}")
        return False
    except Exception as e:
        print(f"Error while opening input document: {e}")
        return False

    print("3. Counting translatable items...")
    total_items = count_translatable_items_in_story_part(doc)
    for story_part in iter_header_footer_parts(doc):
        total_items += count_translatable_items_in_story_part(story_part)

    print(f"   Items to translate: {total_items}")
    progress = TranslationProgress(total_items)
    if total_items > 0:
        progress.render()
    else:
        print("   No translatable items found. The document will be saved unchanged.")
    
    print("\n4. Translating body content...")
    translate_story_part(
        doc,
        context,
        ollama_url,
        model_context,
        source_language,
        target_language,
        progress=progress,
    )

    print("5. Translating headers and footers...")
    for story_part in iter_header_footer_parts(doc):
        translate_story_part(
            story_part,
            context,
            ollama_url,
            model_context,
            source_language,
            target_language,
            progress=progress,
        )

    progress.finish()

    print("6. Saving the translated document...")
    try:
        doc.save(output_path)
    except PermissionError as e:
        print(f"Error: translated output file cannot be generated (permission denied): {output_path}")
        print(f"Details: {e}")
        return False
    except OSError as e:
        print(f"Error: translated output file cannot be generated: {output_path}")
        print(f"Details: {e}")
        return False

    print(f"Translation complete! File saved as: {output_path}")

    print("7. Unloading the model from memory...")
    free_ollama_memory(ollama_url=ollama_url, model=model_context)
    return True

if __name__ == "__main__":
    success = translate_full_docx(INPUT_DOCX_PATH, OUTPUT_DOCX_PATH, OLLAMA_URL, MODEL_CONTEXT, LANGUAGE_SOURCE, LANGUAGE_TARGET)
    if not success:
        print("Translation aborted due to one or more errors.")